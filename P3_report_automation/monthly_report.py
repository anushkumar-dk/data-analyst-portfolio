"""Merges every CSV in ./data, cleans it, writes report.xlsx (Summary + Clean Data). Drop new files in and re-run."""
import glob, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
frames=[]
for f in sorted(glob.glob("data/*.csv")):
    d=pd.read_csv(f); d.columns=[c.strip().lower() for c in d.columns]
    d=d.rename(columns={"product name":"product","quantity":"qty","revenue (inr)":"revenue"})
    d["source_file"]=f.split("/")[-1]; frames.append(d)
df=pd.concat(frames,ignore_index=True); n=len(df)
df["product"]=df["product"].str.strip().str.title()
df=df.drop_duplicates(subset=["date","product","region","qty","revenue"])
df["date"]=pd.to_datetime(df["date"]); df["month"]=df["date"].dt.strftime("%Y-%m")
df["qty"]=df["qty"].fillna(df["qty"].median()).astype(int)
df=df[["date","product","region","qty","revenue","month","source_file"]]
print(f"Merged {len(frames)} files, {n} rows -> {len(df)} after cleaning")
A=Font(name="Arial"); AB=Font(name="Arial",bold=True); HD=Font(name="Arial",bold=True,color="FFFFFF"); FILL=PatternFill("solid",fgColor="1F4E78")
wb=Workbook(); ws=wb.active; ws.title="Summary"; cd=wb.create_sheet("Clean Data")
cd.append(list(df.columns))
for row in df.itertuples(index=False): cd.append([row.date.date(),row.product,row.region,row.qty,row.revenue,row.month,row.source_file])
for c in cd[1]: c.font=HD; c.fill=FILL
N=len(df)+1
ws["A1"]="Monthly Sales Summary"; ws["A1"].font=Font(name="Arial",size=14,bold=True)
r=3
for title,col,L in (("By month","month","F"),("By region","region","C"),("By product","product","B")):
    for j,h in enumerate([title,"Units","Revenue (INR)"]):
        c=ws.cell(r,1+j,h); c.font=HD; c.fill=FILL
    first=r+1
    for it in sorted(df[col].unique()):
        r+=1; ws.cell(r,1,it).font=A
        ws.cell(r,2,f"=SUMIFS('Clean Data'!$D$2:$D${N},'Clean Data'!${L}$2:${L}${N},A{r})").font=A
        c=ws.cell(r,3,f"=SUMIFS('Clean Data'!$E$2:$E${N},'Clean Data'!${L}$2:${L}${N},A{r})"); c.number_format="#,##0"; c.font=A
    r+=1; ws.cell(r,1,"Total").font=AB
    ws.cell(r,2,f"=SUM(B{first}:B{r-1})").font=AB
    c=ws.cell(r,3,f"=SUM(C{first}:C{r-1})"); c.number_format="#,##0"; c.font=AB
    r+=2
for k,w in zip("ABC",(22,12,18)): ws.column_dimensions[k].width=w
for k,w in zip("ABCDEFG",(12,14,10,8,12,10,20)): cd.column_dimensions[k].width=w
wb.save("report.xlsx")
